"""Build the portfolio tracking workbooks.

    python templates/portfolio_tracker/build_tracker.py

Writes two files next to this script:
  Portfolio_Tracker_Bloomberg.xlsx  prices, names, currencies, FX, benchmark and risk-free rate pulled with BDP/BDH
  Portfolio_Tracker_Manual.xlsx     same workbook, those cells are inputs instead (works without Bloomberg)

Conventions: blue text = input, black = formula, green = link to another sheet, yellow fill = key setting.
All example rows are marked "Example" and are meant to be replaced.
"""
import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
F = "Arial"
NAVY, INK, MUTED, RULE = "1F3864", "1A1A1A", "5F6670", "D0D5DC"
f_base = Font(name=F, size=10, color=INK)
f_bold = Font(name=F, size=10, bold=True, color=INK)
f_title = Font(name=F, size=18, bold=True, color=NAVY)
f_sub = Font(name=F, size=10, italic=True, color=MUTED)
f_h = Font(name=F, size=10, bold=True, color="FFFFFF")
f_in = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="008000")
f_kpi = Font(name=F, size=16, bold=True, color=NAVY)
f_kpil = Font(name=F, size=9, color=MUTED)
f_sec = Font(name=F, size=12, bold=True, color=NAVY)
fill_h = PatternFill("solid", fgColor=NAVY)
fill_key = PatternFill("solid", fgColor="FFFF00")
fill_card = PatternFill("solid", fgColor="F3F5F8")
fill_ex = PatternFill("solid", fgColor="FFF4E5")
thin = Side(style="thin", color=RULE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
bottom = Border(bottom=Side(style="medium", color=NAVY))
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
PCT, PCT2, EUR0, EUR2, NUM2, DATE = '0.0%;-0.0%;-', '0.00%;-0.00%;-', '#,##0 "€";-#,##0 "€";-', '#,##0.00 "€";-#,##0.00 "€";-', '#,##0.00;-#,##0.00;-', 'dd/mm/yyyy'

H0, H1 = 8, 47          # Holdings data rows (40 lines)
T0, T1 = 6, 505         # Transactions rows (500 lines)
Y0, Y1 = 6, 185         # History rows (180 months = 15 years)
HRANGE = lambda c: f"Holdings!${c}${H0}:${c}${H1}"
TR = lambda c: f"Transactions!${c}${T0}:${c}${T1}"
YR = lambda c: f"History!${c}${Y0}:${c}${Y1}"

ASSET_CLASSES = ["Equity", "Bond", "Commodity", "Money market", "Other"]
STRATEGIES = [("Value strategy", 0.56), ("Trend ETF", 0.24), ("Cash & money market", 0.20)]
REGIONS = ["Global", "North America", "Europe", "Asia-Pacific", "Emerging markets", "Other"]
CURRENCIES = ["EUR", "USD", "GBP", "CHF", "JPY", "Other"]
TX_TYPES = ["Deposit", "Withdrawal", "Buy", "Sell", "Dividend", "Fee"]

EX_HOLD = [  # ticker, name, asset class, strategy, region, ccy, example price, example fx, target weight
    ("IWDA LN Equity", "iShares Core MSCI World UCITS ETF (Acc)", "Equity", "Value strategy", "Global", "USD", 112.40, 1.133, 0.40),
    ("EXS1 GY Equity", "iShares Core DAX UCITS ETF (DE)", "Equity", "Value strategy", "Europe", "EUR", 205.10, 1.0, 0.16),
    ("SGLN LN Equity", "iShares Physical Gold ETC", "Commodity", "Trend ETF", "Global", "GBp", 6180.0, 0.854, 0.12),
    ("IBTM LN Equity", "iShares $ Treasury Bond 7-10yr UCITS ETF", "Bond", "Trend ETF", "North America", "USD", 168.20, 1.133, 0.12),
    ("XEON GY Equity", "Xtrackers II EUR Overnight Rate Swap UCITS ETF", "Money market", "Cash & money market", "Europe", "EUR", 149.30, 1.0, 0.15),
]
EX_TX = [  # date, type, ticker, qty, price, ccy, fx, fees, cash amount, comment
    (dt.date(2025, 1, 31), "Deposit", "", None, None, "EUR", None, 0, 10000, "Initial deposit"),
    (dt.date(2025, 2, 3), "Buy", "IWDA LN Equity", 35, 101.20, "USD", 1.036, 3, None, "Value sleeve"),
    (dt.date(2025, 2, 3), "Buy", "EXS1 GY Equity", 9, 182.50, "EUR", 1.0, 3, None, "Value sleeve"),
    (dt.date(2025, 2, 3), "Buy", "SGLN LN Equity", 18, 4620.0, "GBp", 0.832, 3, None, "Trend signal: gold long"),
    (dt.date(2025, 2, 3), "Buy", "XEON GY Equity", 12, 144.10, "EUR", 1.0, 2, None, "Cash buffer"),
    (dt.date(2025, 7, 31), "Deposit", "", None, None, "EUR", None, 0, 2000, "Monthly savings"),
    (dt.date(2025, 8, 1), "Buy", "IBTM LN Equity", 9, 160.40, "USD", 1.142, 3, None, "Trend signal: Treasuries long"),
    (dt.date(2025, 8, 1), "Buy", "IWDA LN Equity", 4, 108.90, "USD", 1.142, 3, None, "Rebalance to target"),
    (dt.date(2026, 3, 2), "Sell", "IBTM LN Equity", 3, 164.00, "USD", 1.118, 3, None, "Trend signal weakened"),
    (dt.date(2026, 3, 31), "Fee", "", None, None, "EUR", None, 0, 5, "Platform fee"),
]


def hdr(ws, row, labels, widths=None, start=1):
    for j, h in enumerate(labels):
        c = ws.cell(row, start + j, h)
        c.font, c.fill, c.alignment, c.border = f_h, fill_h, center, box
    ws.row_dimensions[row].height = 32
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[L(start + j)].width = w


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f_title
    ws["A2"] = sub
    ws["A2"].font = f_sub
    ws.sheet_view.showGridLines = False


def style(c, font=f_base, fmt=None, fill=None, border=box):
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = border
    return c


def build(bbg: bool):
    wb = Workbook()
    kind = "Bloomberg" if bbg else "Manual"

    # ======================================================================= Guide
    g = wb.active
    g.title = "Guide"
    title(g, "Portfolio Tracker", f"{kind} version · Strategies Lab · base currency EUR")
    g.column_dimensions["A"].width = 4
    g.column_dimensions["B"].width = 30
    g.column_dimensions["C"].width = 100
    rows = [
        ("What it does", "Tracks a multi-asset portfolio in EUR: positions and P&L, allocation vs targets, monthly performance vs a benchmark, risk metrics (volatility, drawdown, VaR, Expected Shortfall, beta, tracking error) and a decision journal."),
        ("Colour code", "Blue text = input you type. Black text = formula, do not overwrite. Green text = link to another sheet. Yellow fill = key setting. Light orange rows = example data to replace."),
        ("1 · Settings", "Set the portfolio name, benchmark, risk-free rate, VaR confidence level, rebalancing threshold and target weight per strategy."),
        ("2 · Transactions", "Log every deposit, withdrawal, buy, sell, dividend and fee, one row each. Quantities and cash are derived from this sheet only."),
        ("3 · Holdings", "List each instrument once (ticker in column A). Quantity, cost, value, weight, P&L and the rebalancing action are calculated."
         + (" Name, currency, last price, FX, 1-day and YTD change come from Bloomberg (BDP)." if bbg else " Type the name, quote currency, last price and FX rate (units of currency per 1 EUR).")),
        ("4 · History", "At each month end, type the date and the portfolio's total value (Dashboard, cell B6). External flows come from Transactions; returns use the Modified Dietz method."
         + (" The benchmark level is pulled with BDH." if bbg else " Type the benchmark level at the same date.")),
        ("5 · Read", "Dashboard: summary and charts. Risk: every ratio with what it means. Allocation: actual vs target by asset class, strategy, region and currency. Journal: decisions and reasons."),
        ("Monthly routine", "1) Log the month's transactions. 2) Refresh prices" + (" (Bloomberg ribbon > Refresh Workbook)." if bbg else " (type last prices and FX).") +
         " 3) Copy the total value into History. 4) Check the Action column in Holdings and the deviations in Allocation. 5) Write one line in the Journal."),
        ("Method notes", "Monthly return (Modified Dietz) = (V_end − V_start − Flow) / (V_start + 0.5 × Flow). Annualised volatility = monthly standard deviation × √12. Sharpe = mean monthly excess return × 12 / annualised volatility. "
         "Historical VaR = loss not exceeded in (confidence)% of months; ES = average loss beyond the VaR. Risk metrics need at least 12 months of history to be meaningful."),
        ("Limitations", "Average cost uses buys only (realised P&L on sells is not tracked separately). FX for cost is the rate typed on each transaction. Monthly data understate intra-month drawdowns."),
    ]
    if bbg:
        rows.append(("Bloomberg", "Open on a terminal with the Excel add-in. If #NAME? appears: Ctrl+Alt+F9. GBp (pence) quotes are converted automatically. Each refresh uses data quota: keep the ticker list tidy."))
    for i, (k, v) in enumerate(rows):
        r = 4 + i
        style(g.cell(r, 2, k), f_bold, border=None).alignment = wrap
        style(g.cell(r, 3, v), border=None).alignment = wrap
        g.row_dimensions[r].height = 46 if len(v) > 140 else 32
    r = 5 + len(rows)
    g.cell(r, 2, "This workbook is a tool, not investment advice. Example figures are illustrative only.").font = f_sub

    # ======================================================================= Settings
    s = wb.create_sheet("Settings")
    title(s, "Settings", "Key parameters. Everything else in the workbook reads from here.")
    s.column_dimensions["A"].width = 34
    s.column_dimensions["B"].width = 26
    s.column_dimensions["C"].width = 70
    settings = [
        (4, "Portfolio name", "Long-term portfolio (example)", None, "Shown on the Dashboard."),
        (5, "Base currency", "EUR", None, "All values are converted to this currency."),
        (6, "Benchmark name", "MSCI World Net TR (EUR)", None, "Used for beta, tracking error, information ratio and charts."),
        (7, "Benchmark Bloomberg ticker", "MSDEWIN Index", None, "Used by BDH in the History sheet (Bloomberg version)."),
        (8, "Risk-free rate (annual)", "=BDP(\"EUR003M Index\",\"PX_LAST\")/100" if bbg else 0.026, PCT2,
         "3-month Euribor" + (" pulled from Bloomberg." if bbg else ", typed by hand. Update it a few times a year.")),
        (9, "VaR / ES confidence level", 0.95, PCT, "0.95 or 0.99."),
        (10, "Rebalancing threshold", 0.02, PCT, "A position is flagged when its weight differs from target by more than this."),
        (11, "Inception date", dt.date(2025, 1, 31), DATE, "First month-end in History."),
    ]
    for r, k, v, fmt, note in settings:
        style(s.cell(r, 1, k), f_bold, border=None)
        c = style(s.cell(r, 2, v), f_in, fmt, fill_key)
        style(s.cell(r, 3, note), f_sub, border=None)
    s.cell(13, 1, "Target weights by strategy").font = f_sec
    hdr(s, 14, ["Strategy", "Target weight", "Note"])
    for i, (name, w) in enumerate(STRATEGIES):
        style(s.cell(15 + i, 1, name), f_in)
        style(s.cell(15 + i, 2, w), f_in, PCT, fill_key)
        style(s.cell(15 + i, 3, "Example: 80 % risky assets split 70/30 value/trend, 20 % cash (IPS)."), f_sub, border=None)
    rt = 15 + len(STRATEGIES)
    style(s.cell(rt, 1, "Total"), f_bold)
    style(s.cell(rt, 2, f"=SUM(B15:B{rt - 1})"), f_bold, PCT)
    style(s.cell(rt, 3, f'=IF(ABS(B{rt}-1)>0.0001,"Targets must add up to 100 %","OK")'), f_sub, border=None)
    lists = [("Asset classes", ASSET_CLASSES), ("Regions", REGIONS), ("Currencies", CURRENCIES), ("Transaction types", TX_TYPES)]
    col = 5
    for name, items in lists:
        style(s.cell(13, col, name), f_sec, border=None)
        s.column_dimensions[L(col)].width = 20
        for i, it in enumerate(items):
            style(s.cell(14 + i, col, it), f_in)
        col += 1
    strat_rng = f"Settings!$A$15:$A${rt - 1}"

    # ======================================================================= Transactions
    t = wb.create_sheet("Transactions")
    title(t, "Transactions", "One row per event. Buy/Sell amounts are computed from quantity × price ÷ FX. For deposits, withdrawals, dividends and fees type the EUR amount in column J.")
    t.freeze_panes = f"A{T0}"
    hdr(t, T0 - 1, ["Date", "Type", "Ticker", "Quantity", "Price (quote ccy)", "Quote currency", "FX (ccy per 1 EUR)",
                    "Fees (EUR)", "Gross amount (EUR)", "Cash amount (EUR)", "Cash impact (EUR)", "Comment"],
        [12, 12, 18, 10, 14, 11, 12, 11, 15, 15, 15, 40])
    dv_type = DataValidation(type="list", formula1=f"=Settings!$H$14:$H${13 + len(TX_TYPES)}", allow_blank=True)
    t.add_data_validation(dv_type)
    dv_type.add(f"B{T0}:B{T1}")
    for r in range(T0, T1 + 1):
        for cidx, fmt in [(1, DATE), (2, None), (3, None), (4, '#,##0.####'), (5, NUM2), (6, None), (7, '0.0000'), (8, EUR2), (10, EUR2), (12, None)]:
            style(t.cell(r, cidx), f_in, fmt)
        style(t.cell(r, 9, f'=IF(OR(B{r}="Buy",B{r}="Sell"),D{r}*E{r}*IF(EXACT(F{r},"GBp"),0.01,1)/IF(G{r}="",1,G{r}),"")'), fmt=EUR2)
        style(t.cell(r, 11, f'=IF(B{r}="","",IF(B{r}="Deposit",J{r},IF(B{r}="Withdrawal",-J{r},IF(B{r}="Buy",-(I{r}+H{r}),'
                            f'IF(B{r}="Sell",I{r}-H{r},IF(B{r}="Dividend",J{r},IF(B{r}="Fee",-J{r},0)))))))'), fmt=EUR2)
    for i, row in enumerate(EX_TX):
        r = T0 + i
        for cidx, v in zip([1, 2, 3, 4, 5, 6, 7, 8, 10, 12], row):
            t.cell(r, cidx).value = v
        for cidx in range(1, 13):
            t.cell(r, cidx).fill = fill_ex

    # ======================================================================= Holdings
    h = wb.create_sheet("Holdings")
    title(h, "Holdings", "One row per instrument. Type the ticker, classification and target weight; the rest is calculated from Transactions"
          + (" and Bloomberg." if bbg else " and the prices you type."))
    h.freeze_panes = f"B{H0}"
    labels = ["Ticker", "Name", "Asset class", "Strategy", "Region", "Quote currency", "Quantity", "Avg cost (EUR / unit)",
              "Last price (quote ccy)", "FX (ccy per 1 EUR)", "Market value (EUR)", "Weight", "Cost basis (EUR)",
              "Unrealised P&L (EUR)", "P&L %", "1-day change", "YTD change", "Target weight", "Deviation", "Action", "Comment"]
    hdr(h, H0 - 1, labels, [18, 34, 13, 18, 15, 10, 10, 13, 13, 11, 15, 9, 14, 14, 9, 9, 9, 10, 10, 24, 30])
    for key, rng, col in [("ac", f"=Settings!$E$14:$E${13 + len(ASSET_CLASSES)}", "C"), ("st", f"={strat_rng}", "D"),
                          ("rg", f"=Settings!$F$14:$F${13 + len(REGIONS)}", "E")]:
        dv = DataValidation(type="list", formula1=rng, allow_blank=True)
        h.add_data_validation(dv)
        dv.add(f"{col}{H0}:{col}{H1}")
    # summary block
    sb = [("Invested value", f"=SUM(K{H0}:K{H1})", EUR0), ("Cash (from Transactions)", f"=SUM({TR('K')})", EUR0),
          ("Total portfolio value", "=B3+B4", EUR0), ("Number of positions", f"=COUNTIF(G{H0}:G{H1},\">0\")", "0")]
    for i, (k, v, fmt) in enumerate(sb):
        style(h.cell(3 + i, 1, k), f_bold, border=None)
        style(h.cell(3 + i, 2, v), f_bold if i == 2 else f_base, fmt, border=None)
    h["D3"] = "Rows with a ticker but zero quantity are kept and shown as closed."
    h["D3"].font = f_sub
    for r in range(H0, H1 + 1):
        A = f"A{r}"
        blank = f'IF({A}="","",'
        if bbg:
            vals = {2: f'={blank}BDP({A},"NAME"))', 6: f'={blank}BDP({A},"CRNCY"))',
                    9: f'={blank}BDP({A},"PX_LAST"))',
                    10: f'={blank}IF(F{r}="EUR",1,BDP("EUR"&IF(EXACT(F{r},"GBp"),"GBP",UPPER(F{r}))&" Curncy","PX_LAST")))',
                    16: f'={blank}BDP({A},"CHG_PCT_1D")/100)', 17: f'={blank}BDP({A},"CHG_PCT_YTD")/100)'}
        else:
            vals = {}
        q = f'SUMIFS({TR("D")},{TR("C")},{A},{TR("B")},"Buy")-SUMIFS({TR("D")},{TR("C")},{A},{TR("B")},"Sell")'
        form = {
            7: f'={blank}{q})',
            8: f'={blank}IFERROR((SUMIFS({TR("I")},{TR("C")},{A},{TR("B")},"Buy")+SUMIFS({TR("H")},{TR("C")},{A},{TR("B")},"Buy"))/SUMIFS({TR("D")},{TR("C")},{A},{TR("B")},"Buy"),0))',
            11: f'={blank}IFERROR(G{r}*I{r}*IF(EXACT(F{r},"GBp"),0.01,1)/J{r},0))',
            12: f'={blank}IFERROR(K{r}/$B$5,0))',
            13: f'={blank}G{r}*H{r})',
            14: f'={blank}K{r}-M{r})',
            15: f'={blank}IFERROR(N{r}/M{r},0))',
            19: f'={blank}L{r}-R{r})',
            20: f'={blank}IF(ABS(S{r})<=Settings!$B$10,"OK",IF(S{r}>0,"Sell ","Buy ")&TEXT(ABS(S{r})*$B$5,"#,##0")&" €"))',
        }
        fmts = {7: '#,##0.####', 8: EUR2, 9: NUM2, 10: '0.0000', 11: EUR0, 12: PCT, 13: EUR0, 14: EUR0, 15: PCT,
                16: PCT2, 17: PCT, 18: PCT, 19: PCT}
        for cidx in range(1, 22):
            if cidx in form:
                style(h.cell(r, cidx, form[cidx]), fmt=fmts.get(cidx))
            elif cidx in vals:
                style(h.cell(r, cidx, vals[cidx]), fmt=fmts.get(cidx))
            else:
                style(h.cell(r, cidx), f_in, fmts.get(cidx))
    for i, ex in enumerate(EX_HOLD):
        r = H0 + i
        tk, name, ac, stg, reg, ccy, px, fx, tw = ex
        h.cell(r, 1).value, h.cell(r, 3).value, h.cell(r, 4).value, h.cell(r, 5).value, h.cell(r, 18).value = tk, ac, stg, reg, tw
        h.cell(r, 21).value = "Example — replace"
        if not bbg:
            h.cell(r, 2).value, h.cell(r, 6).value, h.cell(r, 9).value, h.cell(r, 10).value = name, ccy, px, fx
        for cidx in range(1, 22):
            h.cell(r, cidx).fill = fill_ex
    h.conditional_formatting.add(f"N{H0}:O{H1}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, size=10, color="B3261E")))
    h.conditional_formatting.add(f"N{H0}:O{H1}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=F, size=10, color="0B6B2F")))
    h.conditional_formatting.add(f"T{H0}:T{H1}", FormulaRule(formula=[f'AND(T{H0}<>"",T{H0}<>"OK")'], fill=PatternFill("solid", fgColor="FDE2C4")))
    style(h.cell(H1 + 1, 1, "Total"), f_bold)
    for cidx, f_ in [(11, f"=SUM(K{H0}:K{H1})"), (12, f"=SUM(L{H0}:L{H1})"), (13, f"=SUM(M{H0}:M{H1})"), (14, f"=SUM(N{H0}:N{H1})"),
                     (15, f"=IFERROR(N{H1 + 1}/M{H1 + 1},0)"), (18, f"=SUM(R{H0}:R{H1})")]:
        style(h.cell(H1 + 1, cidx, f_), f_bold, fmts.get(cidx, EUR0))
    style(h.cell(H1 + 2, 1, "Cash"), f_bold)
    style(h.cell(H1 + 2, 11, "=B4"), f_bold, EUR0)
    style(h.cell(H1 + 2, 12, "=IFERROR(B4/B5,0)"), f_bold, PCT)

    # ======================================================================= History
    y = wb.create_sheet("History")
    title(y, "History", "One row per month end. Type the date and total value (and the benchmark level"
          + (" is pulled by BDH)." if bbg else " in column E)."))
    y.freeze_panes = f"B{Y0}"
    hdr(y, Y0 - 1, ["Month end", "Portfolio value (EUR)", "Net external flows (EUR)", "Portfolio return", "Benchmark level",
                    "Benchmark return", "Risk-free (monthly)", "Excess return", "Benchmark excess", "Active return",
                    "NAV portfolio (base 100)", "NAV benchmark (base 100)", "Drawdown", "Benchmark drawdown",
                    "Downside² (helper)", "Year (helper)", "ln(1+r) (helper)", "Comment"],
        [12, 15, 15, 11, 13, 11, 11, 11, 11, 11, 13, 13, 11, 11, 11, 9, 11, 34])
    for r in range(Y0, Y1 + 1):
        A, p = f"A{r}", r - 1
        first = r == Y0
        b = f'IF(OR({A}="",B{r}=""),"",'
        flows = (f'=IF({A}="","",SUMIFS({TR("J")},{TR("B")},"Deposit",{TR("A")},">"&{"A" + str(p) if not first else "0"},{TR("A")},"<="&{A})'
                 f'-SUMIFS({TR("J")},{TR("B")},"Withdrawal",{TR("A")},">"&{"A" + str(p) if not first else "0"},{TR("A")},"<="&{A}))')
        row = {
            3: flows,
            4: '=""' if first else f'={b}IF(B{p}="","",(B{r}-B{p}-C{r})/(B{p}+0.5*C{r})))',
            5: (f'=IF({A}="","",BDH(Settings!$B$7,"PX_LAST",{A},{A},"Days=A","Fill=P","Dts=H"))' if bbg else None),
            6: '=""' if first else f'=IF(OR(E{r}="",E{p}=""),"",E{r}/E{p}-1)',
            7: f'=IF({A}="","",Settings!$B$8/12)',
            8: f'=IF(D{r}="","",D{r}-G{r})',
            9: f'=IF(F{r}="","",F{r}-G{r})',
            10: f'=IF(OR(D{r}="",F{r}=""),"",D{r}-F{r})',
            11: f'={b}100)' if first else f'={b}IF(D{r}="",K{p},K{p}*(1+D{r})))',
            12: f'=IF(E{r}="","",100)' if first else f'=IF(E{r}="","",IF(F{r}="",L{p},L{p}*(1+F{r})))',
            13: f'=IF(K{r}="","",K{r}/MAX(K${Y0}:K{r})-1)',
            14: f'=IF(L{r}="","",L{r}/MAX(L${Y0}:L{r})-1)',
            15: f'=IF(H{r}="","",MIN(H{r},0)^2)',
            16: f'=IF({A}="","",YEAR({A}))',
            17: f'=IF(D{r}="","",LN(1+D{r}))',
        }
        fm = {1: DATE, 2: EUR0, 3: EUR0, 4: PCT2, 5: NUM2, 6: PCT2, 7: PCT2, 8: PCT2, 9: PCT2, 10: PCT2, 11: '0.0', 12: '0.0',
              13: PCT, 14: PCT, 15: '0.000000', 16: '0', 17: '0.0000'}
        for cidx in range(1, 19):
            v = row.get(cidx)
            if v is None:
                style(y.cell(r, cidx), f_in, fm.get(cidx))
            else:
                style(y.cell(r, cidx, v), f_link if cidx == 7 else f_base, fm.get(cidx))
    # example history: 21 month ends, consistent with example flows
    random.seed(7)
    months = []
    d = dt.date(2025, 1, 31)
    while d <= dt.date(2026, 9, 30):
        months.append(d)
        nm = dt.date(d.year + (d.month // 12), d.month % 12 + 1, 1)
        d = dt.date(nm.year + (nm.month // 12), nm.month % 12 + 1, 1) - dt.timedelta(days=1)
    v, bench = 10000.0, 412.0
    for i, m in enumerate(months):
        r = Y0 + i
        if i:
            rb = random.gauss(0.008, 0.035)
            rp = 0.7 * rb + random.gauss(0.002, 0.012)
            flow = 2000 if m == dt.date(2025, 7, 31) else 0
            v = v * (1 + rp) + flow
            bench *= 1 + rb
        y.cell(r, 1).value = m
        y.cell(r, 2).value = round(v, 2) if i < len(months) - 1 else 12679.23  # last month = example Holdings total
        if not bbg:
            y.cell(r, 5).value = round(bench, 2)
        y.cell(r, 18).value = "Example — replace"
        for cidx in range(1, 19):
            y.cell(r, cidx).fill = fill_ex

    # ======================================================================= Risk
    k = wb.create_sheet("Risk")
    title(k, "Risk & ratios", "Monthly data from History. Each figure is compared with the benchmark and explained in plain words.")
    hdr(k, 4, ["Metric", "Portfolio", "Benchmark", "Reading", "What it means"], [32, 14, 14, 16, 90])
    D_, F_, H_, I_, J_, M_, N_, O_ = (YR(c) for c in "DFHIJMNO")
    last_nav = f"LOOKUP(9.99E+307,{YR('K')})"
    last_navb = f"LOOKUP(9.99E+307,{YR('L')})"
    conf = "Settings!$B$9"
    metrics = [
        ("Months of history", f"=COUNT({D_})", f"=COUNT({F_})", None, "Number of monthly returns used. Below 12 months, treat every ratio as indicative only.", "0"),
        ("Total return (time-weighted)", f"=IFERROR({last_nav}/100-1,0)", f"=IFERROR({last_navb}/100-1,0)", None,
         "Growth of 100 invested at inception, excluding the effect of deposits and withdrawals.", PCT),
        ("Annualised return", "=IFERROR((1+B6)^(12/B5)-1,0)", "=IFERROR((1+C6)^(12/C5)-1,0)", None, "Compound yearly growth rate over the history.", PCT),
        ("Year-to-date return", f"=IFERROR(EXP(SUMIFS({YR('Q')},{YR('P')},YEAR(LOOKUP(9.99E+307,{YR('A')}))))-1,0)", None, None,
         "Time-weighted return since 1 January of the current year.", PCT),
        ("Annualised volatility", f"=IFERROR(STDEV({D_})*SQRT(12),0)", f"=IFERROR(STDEV({F_})*SQRT(12),0)", None,
         "Typical size of yearly fluctuations. Monthly standard deviation × √12.", PCT),
        ("Sharpe ratio", f"=IFERROR(AVERAGE({H_})*12/(STDEV({H_})*SQRT(12)),0)", f"=IFERROR(AVERAGE({I_})*12/(STDEV({I_})*SQRT(12)),0)",
         'SH', "Excess return per unit of total risk. Above 0.5 is good for a diversified portfolio, above 1 is rare over long periods.", NUM2),
        ("Sortino ratio", f"=IFERROR(AVERAGE({H_})*12/(SQRT(SUM({O_})/COUNT({O_}))*SQRT(12)),0)", None, None,
         "Like Sharpe, but only penalises downside volatility.", NUM2),
        ("Maximum drawdown", f"=IFERROR(MIN({M_}),0)", f"=IFERROR(MIN({N_}),0)", None,
         "Largest fall from a peak to a later trough (month-end data, so intra-month falls can be deeper).", PCT),
        ("Current drawdown", f"=IFERROR(LOOKUP(2,1/({YR('M')}<>\"\"),{YR('M')}),0)", None, None, "Distance from the highest value reached so far.", PCT),
        ("Calmar ratio", "=IFERROR(B7/ABS(B12),0)", "=IFERROR(C7/ABS(C12),0)", None, "Annualised return divided by the maximum drawdown.", NUM2),
        ("Historical VaR (monthly)", f"=IFERROR(-PERCENTILE({D_},1-{conf}),0)", f"=IFERROR(-PERCENTILE({F_},1-{conf}),0)", None,
         "Monthly loss not exceeded in the chosen share of months (Settings!B9). A 95 % VaR of 5 % means one month in twenty loses more than 5 %.", PCT2),
        ("Expected Shortfall (monthly)", f"=IFERROR(-AVERAGEIF({D_},\"<=\"&-B15),0)", f"=IFERROR(-AVERAGEIF({F_},\"<=\"&-C15),0)", None,
         "Average loss in the months worse than the VaR. Captures how deep the bad months are.", PCT2),
        ("Parametric VaR (monthly, normal)", f"=IFERROR(-(AVERAGE({D_})+NORMSINV(1-{conf})*STDEV({D_})),0)", None, None,
         "VaR assuming normally distributed returns. If it is below the historical VaR, losses have fatter tails than a normal law.", PCT2),
        ("VaR in euros (monthly)", "=B15*Dashboard!$B$6", None, None, "Historical VaR applied to today's portfolio value.", EUR0),
        ("ES in euros (monthly)", "=B16*Dashboard!$B$6", None, None, "Expected Shortfall applied to today's portfolio value.", EUR0),
        ("Beta vs benchmark", f"=IFERROR(SLOPE({D_},{F_}),0)", "1", None, "Sensitivity to the benchmark: 0.7 means the portfolio moves 0.7 % when the benchmark moves 1 %.", NUM2),
        ("Correlation vs benchmark", f"=IFERROR(CORREL({D_},{F_}),0)", "1", None, "From −1 (opposite) to +1 (identical). Lower means more diversification.", NUM2),
        ("Jensen's alpha (annual)", f"=IFERROR((AVERAGE({H_})-B20*AVERAGE({I_}))*12,0)", None, None, "Return above what the beta alone would explain.", PCT2),
        ("Tracking error", f"=IFERROR(STDEV({J_})*SQRT(12),0)", None, None, "Volatility of the difference with the benchmark. How far the portfolio strays from it.", PCT),
        ("Information ratio", f"=IFERROR(AVERAGE({J_})*12/B23,0)", None, None, "Active return per unit of tracking error. Above 0.5 is good.", NUM2),
        ("Up-capture", f"=IFERROR(AVERAGEIF({F_},\">0\",{D_})/AVERAGEIF({F_},\">0\"),0)", None, None, "Share of the benchmark's rising months captured.", PCT),
        ("Down-capture", f"=IFERROR(AVERAGEIF({F_},\"<0\",{D_})/AVERAGEIF({F_},\"<0\"),0)", None, None, "Share of the benchmark's falling months suffered. Below 100 % is what diversification aims for.", PCT),
        ("Positive months", f"=IFERROR(COUNTIF({D_},\">0\")/COUNT({D_}),0)", f"=IFERROR(COUNTIF({F_},\">0\")/COUNT({F_}),0)", None, "Share of months with a gain.", PCT),
        ("Best month", f"=IFERROR(MAX({D_}),0)", f"=IFERROR(MAX({F_}),0)", None, "", PCT2),
        ("Worst month", f"=IFERROR(MIN({D_}),0)", f"=IFERROR(MIN({F_}),0)", None, "", PCT2),
        ("Skewness", f"=IFERROR(SKEW({D_}),0)", f"=IFERROR(SKEW({F_}),0)", None, "Negative = more large losses than large gains (typical of equities).", NUM2),
        ("Excess kurtosis", f"=IFERROR(KURT({D_}),0)", f"=IFERROR(KURT({F_}),0)", None, "Above 0 = fatter tails than a normal distribution.", NUM2),
        ("Concentration (HHI)", f"=SUMSQ({HRANGE('L')})+Holdings!$L${H1 + 2}^2", None, None, "Sum of squared weights. 1/HHI = effective number of positions.", '0.000'),
        ("Effective number of positions", "=IFERROR(1/B32,0)", None, None, "Number of equally weighted positions with the same concentration.", '0.0'),
        ("Top-3 weight", f"=IFERROR(SUMPRODUCT(LARGE({HRANGE('L')},{{1,2,3}})),SUM({HRANGE('L')}))", None, None, "Weight of the three largest positions.", PCT),
        ("Cash weight", f"=Holdings!$L${H1 + 2}", None, None, "Share of the portfolio held in cash (money-market ETFs count as positions).", PCT),
        ("Fees paid (total)", f"=SUM({TR('H')})+SUMIFS({TR('J')},{TR('B')},\"Fee\")", None, None, "Transaction fees plus platform fees since inception.", EUR0),
        ("Dividends received (total)", f"=SUMIFS({TR('J')},{TR('B')},\"Dividend\")", None, None, "Cash dividends logged in Transactions.", EUR0),
    ]
    for i, (name, fp, fb, reading, note, fmt) in enumerate(metrics):
        r = 5 + i
        style(k.cell(r, 1, name), f_bold)
        style(k.cell(r, 2, fp), fmt=fmt)
        style(k.cell(r, 3, fb if fb is not None else ""), fmt=fmt)
        if reading == "SH":
            rd = f'=IF(B{r}>=1,"Strong",IF(B{r}>=0.5,"Good",IF(B{r}>=0,"Weak","Negative")))'
        elif name == "Historical VaR (monthly)":
            rd = f'=IF(B{r}>C{r},"Riskier than benchmark","Less risky than benchmark")'
        elif name == "Maximum drawdown":
            rd = f'=IF(B{r}<C{r},"Deeper than benchmark","Shallower than benchmark")'
        elif name == "Correlation vs benchmark":
            rd = f'=IF(B{r}>0.9,"Very close to index",IF(B{r}>0.6,"Moderately diversified","Well diversified"))'
        elif name == "Annualised volatility":
            rd = f'=IF(B{r}>C{r},"More volatile","Less volatile")'
        else:
            rd = ""
        style(k.cell(r, 4, rd), f_sub)
        c = style(k.cell(r, 5, note), f_sub)
        c.alignment = wrap
        k.row_dimensions[r].height = 28
    k.freeze_panes = "B5"

    # ======================================================================= Allocation
    a = wb.create_sheet("Allocation")
    title(a, "Allocation", "Actual weights from Holdings (cash included) compared with targets.")
    a.column_dimensions["A"].width = 24
    for c_ in "BCDE":
        a.column_dimensions[c_].width = 14
    tot = "Holdings!$B$5"
    cash = "Holdings!$B$4"

    def block(r0, label, items, col, extra_cash=None, targets=None):
        style(a.cell(r0, 1, label), f_sec, border=None)
        heads = ["Group", "Value (EUR)", "Weight"] + (["Target", "Deviation"] if targets else [])
        hdr(a, r0 + 1, heads)
        for i, it in enumerate(items):
            r = r0 + 2 + i
            style(a.cell(r, 1, it if not targets else f"=Settings!A{15 + i}"), f_link if targets else f_base)
            v = f"=SUMIFS({HRANGE('K')},{HRANGE(col)},A{r})"
            if extra_cash and it == extra_cash:
                v += f"+{cash}"
            style(a.cell(r, 2, v), fmt=EUR0)
            style(a.cell(r, 3, f"=IFERROR(B{r}/{tot},0)"), fmt=PCT)
            if targets:
                style(a.cell(r, 4, f"=Settings!B{15 + i}"), f_link, PCT)
                style(a.cell(r, 5, f"=C{r}-D{r}"), fmt=PCT)
        rT = r0 + 2 + len(items)
        style(a.cell(rT, 1, "Total"), f_bold)
        style(a.cell(rT, 2, f"=SUM(B{r0 + 2}:B{rT - 1})"), f_bold, EUR0)
        style(a.cell(rT, 3, f"=SUM(C{r0 + 2}:C{rT - 1})"), f_bold, PCT)
        return r0 + 2, rT - 1

    ac0, ac1 = block(4, "By asset class", ASSET_CLASSES + ["Cash"], "C", extra_cash="Cash")
    a.cell(ac1, 2).value = f"={cash}"
    st0, st1 = block(ac1 + 3, "By strategy", [n for n, _ in STRATEGIES], "D", extra_cash="Cash & money market", targets=True)
    a.cell(st1, 2).value = f"=SUMIFS({HRANGE('K')},{HRANGE('D')},A{st1})+{cash}"
    rg0, rg1 = block(st1 + 3, "By region", REGIONS, "E")
    a.conditional_formatting.add(f"E{st0}:E{st1}", FormulaRule(formula=[f"ABS(E{st0})>Settings!$B$10"], fill=PatternFill("solid", fgColor="FDE2C4")))
    # currency exposure: GBp counted as GBP
    cu_r = rg1 + 3
    style(a.cell(cu_r, 1, "By currency (before hedging)"), f_sec, border=None)
    hdr(a, cu_r + 1, ["Currency", "Value (EUR)", "Weight"])
    for i, cur in enumerate(CURRENCIES[:-1]):
        r = cu_r + 2 + i
        style(a.cell(r, 1, cur))
        v = f"=SUMIFS({HRANGE('K')},{HRANGE('F')},A{r})"
        if cur == "GBP":
            v = f'=SUMPRODUCT(({HRANGE("F")}="GBP")+EXACT({HRANGE("F")},"GBp"),{HRANGE("K")})'
        if cur == "EUR":
            v += f"+{cash}"
        style(a.cell(r, 2, v), fmt=EUR0)
        style(a.cell(r, 3, f"=IFERROR(B{r}/{tot},0)"), fmt=PCT)
    cu1 = cu_r + 2 + len(CURRENCIES) - 2

    # charts on Allocation
    pie = PieChart()
    pie.title = "Asset classes"
    pie.add_data(Reference(a, min_col=2, min_row=ac0, max_row=ac1), titles_from_data=False)
    pie.set_categories(Reference(a, min_col=1, min_row=ac0, max_row=ac1))
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.height, pie.width = 7.5, 11
    a.add_chart(pie, "G4")
    bar = BarChart()
    bar.type = "col"
    bar.title = "Strategy: actual vs target"
    bar.add_data(Reference(a, min_col=3, max_col=4, min_row=st0 - 1, max_row=st1), titles_from_data=True)
    bar.set_categories(Reference(a, min_col=1, min_row=st0, max_row=st1))
    bar.y_axis.numFmt = '0%'
    bar.height, bar.width = 7.5, 11
    a.add_chart(bar, "G20")
    bc = BarChart()
    bc.type = "bar"
    bc.title = "Currency exposure"
    bc.add_data(Reference(a, min_col=3, min_row=cu_r + 2, max_row=cu1), titles_from_data=False)
    bc.set_categories(Reference(a, min_col=1, min_row=cu_r + 2, max_row=cu1))
    bc.legend = None
    bc.x_axis.numFmt = '0%'
    bc.height, bc.width = 7.5, 11
    a.add_chart(bc, "G36")

    # ======================================================================= Journal
    j = wb.create_sheet("Journal")
    title(j, "Decision journal", "One line per decision or review: what was done, why, and what would prove it wrong.")
    hdr(j, 4, ["Date", "Type", "Decision / observation", "Reasoning", "What would change my mind", "Review date", "Outcome"],
        [12, 14, 40, 50, 40, 12, 30])
    dvj = DataValidation(type="list", formula1='"Rebalance,New position,Exit,Review,Risk alert,Other"', allow_blank=True)
    j.add_data_validation(dvj)
    dvj.add("B5:B300")
    for r in range(5, 301):
        for cidx in range(1, 8):
            c = style(j.cell(r, cidx), f_in, DATE if cidx in (1, 6) else None)
            c.alignment = wrap
    ex = [(dt.date(2025, 2, 3), "New position", "Initial allocation: 56 % value, 24 % trend ETF, 20 % cash.", "IPS: 80/20 risky/cash; trend sleeve diversifies the value sleeve (backtest correlation −0.21 for the futures version).", "Sharpe of the combined portfolio below the value sleeve alone over 3 years.", dt.date(2026, 2, 3), "Example — replace"),
          (dt.date(2026, 3, 2), "Rebalance", "Sold 3 units of the Treasury ETF.", "Trend signal on US 10-year turned negative on two of three horizons.", "Signal back to +1 on all horizons.", dt.date(2026, 6, 30), "Example — replace")]
    for i, row in enumerate(ex):
        for cidx, v in enumerate(row, 1):
            j.cell(5 + i, cidx).value = v
            j.cell(5 + i, cidx).fill = fill_ex
        j.row_dimensions[5 + i].height = 42

    # ======================================================================= ChartData (hidden helper)
    cd = wb.create_sheet("ChartData")
    cd["A1"] = "Helper for charts: #N/A hides empty months in line charts. Do not edit."
    hdr(cd, 3, ["Month end", "Portfolio", "Benchmark", "Drawdown", "Bench. drawdown", "Monthly return"])
    for i, r in enumerate(range(Y0, Y1 + 1)):
        rr = 4 + i
        cd.cell(rr, 1, f'=IF(History!A{r}="",NA(),History!A{r})').number_format = DATE
        cd.cell(rr, 2, f'=IF(History!K{r}="",NA(),History!K{r})').number_format = '0.0'
        cd.cell(rr, 3, f'=IF(History!L{r}="",NA(),History!L{r})').number_format = '0.0'
        cd.cell(rr, 4, f'=IF(History!M{r}="",NA(),History!M{r})').number_format = PCT
        cd.cell(rr, 5, f'=IF(History!N{r}="",NA(),History!N{r})').number_format = PCT
        cd.cell(rr, 6, f'=IF(History!D{r}="",NA(),History!D{r})').number_format = PCT
    # last 24 months window (bar chart of monthly returns)
    cd["K3"], cd["L3"] = "Last 24 months", "Return"
    for i in range(24):
        rr = 4 + i
        pos = f"(IFERROR(MATCH(9.99E+307,{YR('A')}),0)-24+{i + 1})"
        cd.cell(rr, 11, f'=IF({pos}<1,"",TEXT(INDEX({YR("A")},{pos}),"mmm yy"))')
        cd.cell(rr, 12, f'=IF({pos}<1,NA(),IF(INDEX({YR("D")},{pos})="",NA(),INDEX({YR("D")},{pos})))').number_format = PCT
    # top holdings by weight
    cd["H3"], cd["I3"] = "Top holdings", "Weight"
    for i in range(10):
        rr = 4 + i
        cd.cell(rr, 9, f'=IFERROR(LARGE({HRANGE("L")},{i + 1}),NA())').number_format = PCT
        cd.cell(rr, 8, f'=IFERROR(INDEX({HRANGE("A")},MATCH(I{rr},{HRANGE("L")},0)),"")')
    cd.sheet_state = "hidden"
    nrows = Y1 - Y0 + 1

    # ======================================================================= Dashboard
    db = wb.create_sheet("Dashboard", 1)
    title(db, "=Settings!B4", "Summary of value, performance, risk and allocation. All figures update from the other sheets.")
    db["A1"].font = f_title
    db["A3"] = f'="As of "&TEXT(IFERROR(LOOKUP(9.99E+307,{YR("A")}),TODAY()),"dd mmm yyyy")&" · benchmark: "&Settings!B6'
    db["A3"].font = f_sub
    for c_ in range(1, 14):
        db.column_dimensions[L(c_)].width = 13
    db.column_dimensions["A"].width = 2
    tiles = [  # (row, col, label, formula, fmt)
        (5, 2, "Total value", "=Holdings!B5", EUR0),
        (5, 4, "Net deposits", f'=SUMIFS({TR("J")},{TR("B")},"Deposit")-SUMIFS({TR("J")},{TR("B")},"Withdrawal")', EUR0),
        (5, 6, "Gain since inception", "=B6-D6", EUR0),
        (5, 8, "Total return (TWR)", "=Risk!B6", PCT),
        (5, 10, "Year to date", "=Risk!B8", PCT),
        (5, 12, "Annualised return", "=Risk!B7", PCT),
        (9, 2, "Volatility (ann.)", "=Risk!B9", PCT),
        (9, 4, "Sharpe ratio", "=Risk!B10", NUM2),
        (9, 6, "Max drawdown", "=Risk!B12", PCT),
        (9, 8, "VaR monthly", "=Risk!B15", PCT2),
        (9, 10, "Beta vs benchmark", "=Risk!B20", NUM2),
        (9, 12, "Positions / cash", '=TEXT(Holdings!B6,"0")&" / "&TEXT(Risk!B35,"0%")', None),
    ]
    for r, c_, lab, fml, fmt in tiles:
        db.merge_cells(start_row=r, start_column=c_, end_row=r, end_column=c_ + 1)
        db.merge_cells(start_row=r + 1, start_column=c_, end_row=r + 2, end_column=c_ + 1)
        lc = db.cell(r, c_, lab)
        lc.font, lc.fill, lc.alignment = f_kpil, fill_card, Alignment(horizontal="left", vertical="bottom", indent=1)
        vc = db.cell(r + 1, c_, fml)
        vc.font, vc.fill, vc.alignment = f_kpi, fill_card, Alignment(horizontal="left", vertical="center", indent=1)
        if fmt:
            vc.number_format = fmt
        for rr in range(r, r + 3):
            for cc in (c_, c_ + 1):
                db.cell(rr, cc).fill = fill_card
    db["B13"] = "Charts: portfolio vs benchmark, drawdown, monthly returns, allocation and top holdings. Risk details in the Risk sheet."
    db["B13"].font = f_sub

    def cd_ref(col):
        return Reference(cd, min_col=col, min_row=4, max_row=3 + nrows)

    def scatter(title_, cols, names, numfmt, anchor):
        ch = ScatterChart()
        ch.title = title_
        ch.style = 2
        xs = cd_ref(1)
        for col, nm in zip(cols, names):
            se = Series(cd_ref(col), xs, title=nm)
            se.marker.symbol = "none"
            se.smooth = False
            ch.series.append(se)
        ch.x_axis.number_format = "mmm yy"
        ch.x_axis.majorTimeUnit = "months"
        ch.y_axis.number_format = numfmt
        ch.height, ch.width = 8, 17
        db.add_chart(ch, anchor)

    scatter("Growth of 100: portfolio vs benchmark", [2, 3], ["Portfolio", "Benchmark"], "0", "B15")
    scatter("Drawdown", [4, 5], ["Portfolio", "Benchmark"], "0%", "H15")
    mb = BarChart()
    mb.type = "col"
    mb.title = "Monthly returns"
    mb.title = "Monthly returns, last 24 months"
    mb.add_data(Reference(cd, min_col=12, min_row=4, max_row=27), titles_from_data=False)
    mb.set_categories(Reference(cd, min_col=11, min_row=4, max_row=27))
    mb.legend = None
    mb.x_axis.tickLblPos = "low"
    mb.y_axis.numFmt = '0%'
    mb.height, mb.width = 8, 17
    db.add_chart(mb, "B32")
    p2 = PieChart()
    p2.title = "Allocation by strategy"
    p2.add_data(Reference(a, min_col=2, min_row=st0, max_row=st1), titles_from_data=False)
    p2.set_categories(Reference(a, min_col=1, min_row=st0, max_row=st1))
    p2.dataLabels = DataLabelList()
    p2.dataLabels.showPercent = True
    p2.height, p2.width = 8, 17
    db.add_chart(p2, "H32")
    tb = BarChart()
    tb.type = "bar"
    tb.title = "Top holdings by weight"
    tb.add_data(Reference(cd, min_col=9, min_row=4, max_row=13), titles_from_data=False)
    tb.set_categories(Reference(cd, min_col=8, min_row=4, max_row=13))
    tb.legend = None
    tb.x_axis.numFmt = '0%'
    tb.height, tb.width = 8, 17
    db.add_chart(tb, "B49")

    # tab colours and order
    for ws_, colr in [(g, "8A919C"), (db, NAVY), (h, "2A78D6"), (t, "2A78D6"), (y, "2A78D6"), (k, "EB6834"), (a, "1BAF7A"), (j, "6250D6"), (s, "8A919C")]:
        ws_.sheet_properties.tabColor = colr
    wb.move_sheet("Holdings", offset=wb.sheetnames.index("Transactions") - wb.sheetnames.index("Holdings"))
    wb.move_sheet("Settings", offset=wb.sheetnames.index("Journal") - wb.sheetnames.index("Settings"))
    wb.active = 1
    out = HERE / f"Portfolio_Tracker_{kind}.xlsx"
    wb.save(out)
    print("saved", out)
    return out


if __name__ == "__main__":
    build(True)
    build(False)
