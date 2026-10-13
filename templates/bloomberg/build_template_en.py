"""Builds the Bloomberg extraction workbook for the trend-following backtest (English version)."""
import datetime as dt
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule

OUT = sys.argv[1] if len(sys.argv) > 1 else "Extract_Trend_Following_Bloomberg_EN.xlsx"

FIRST = 4          # first data row in Prices
LAST = 7503        # last reserved row (7,500 rows ≈ 28 years of sessions)

F = "Arial"
f_base = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True)
f_h = Font(name=F, size=10, bold=True, color="FFFFFF")
f_input = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="008000")
f_link_b = Font(name=F, size=10, color="008000", bold=True)
f_note = Font(name=F, size=9, italic=True, color="595959")
fill_h = PatternFill("solid", fgColor="1F3864")
fill_in = PatternFill("solid", fgColor="FFFF00")
fill_sub = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

TICKERS = [
    # block, ticker, description, status, is it a future?
    ("Equity futures", "ES1 Index", "S&P 500 E-mini (CME)", "OK", True),
    ("Equity futures", "VG1 Index", "Euro Stoxx 50 (Eurex)", "OK", True),
    ("Equity futures", "NK1 Index", "Nikkei 225 (Osaka)", "OK", True),
    ("Equity futures", "MES1 Index", "MSCI Emerging Markets (ICE)", "OK", True),
    ("Rate futures", "TY1 Comdty", "US Treasury 10-year (CBOT)", "OK", True),
    ("Rate futures", "US1 Comdty", "US Treasury Bond 30-year (CBOT)", "OK", True),
    ("Rate futures", "RX1 Comdty", "Euro-Bund 10-year (Eurex)", "OK", True),
    ("Rate futures", "G 1 Comdty", "Long Gilt (ICE)", "OK", True),
    ("Commodity futures", "GC1 Comdty", "Gold (COMEX)", "OK", True),
    ("Commodity futures", "SI1 Comdty", "Silver (COMEX)", "OK", True),
    ("Commodity futures", "CL1 Comdty", "WTI crude oil (NYMEX) — the backtest uses CL2, see Plan B", "OK", True),
    ("Commodity futures", "HG1 Comdty", "Copper (COMEX)", "OK", True),
    ("Commodity futures", "C 1 Comdty", "Corn (CBOT)", "OK", True),
    ("Currency futures", "EC1 Curncy", "EUR/USD (CME)", "OK", True),
    ("Currency futures", "JY1 Curncy", "JPY/USD (CME)", "OK", True),
    ("Currency futures", "BP1 Curncy", "GBP/USD (CME)", "OK", True),
    ("Currency futures", "AD1 Curncy", "AUD/USD (CME)", "OK", True),
    ("Spot FX", "EURUSD Curncy", "EUR/USD spot", "OK", False),
    ("Spot FX", "EURJPY Curncy", "EUR/JPY spot", "OK", False),
    ("Spot FX", "EURGBP Curncy", "EUR/GBP spot", "OK", False),
    ("Short rates", "USGG3M Index", "US Treasury 3-month (yield, %)", "OK", False),
    ("Short rates", "EUR003M Index", "3-month Euribor (%)", "OK", False),
    ("Equity benchmark", "MSDEWIN Index", "MSCI World Net TR EUR", "OK", False),
    ("Equity benchmark", "M1WO000V Index", "MSCI World Value Net TR (in USD, converted to EUR in the backtest)", "OK", False),
]
def second(tk):
    root, yk = tk.rsplit(" ", 1)
    return root[:-1] + "2 " + yk
TICKERS += [("Plan B — 2nd generics", second(tk), "2nd contract: " + desc, "Plan B", True)
            for bloc, tk, desc, stat, fut in list(TICKERS) if fut]
N = len(TICKERS)

wb = Workbook()

# ---------------------------------------------------------------- Parameters
p = wb.active
p.title = "Parameters"
p.sheet_view.showGridLines = False
p.column_dimensions["A"].width = 34
p.column_dimensions["B"].width = 22
p.column_dimensions["C"].width = 70

p["A1"] = "Bloomberg extraction — multi-asset trend-following backtest"
p["A1"].font = f_title
p["A2"] = "Yellow cells with blue text are the only ones to fill in. Everything else is calculated."
p["A2"].font = f_note

rows = [
    (4, "Start date", dt.date(2000, 1, 3), "dd/mm/yyyy", "First business day of 2000."),
    (5, "End date", dt.date(2026, 9, 30), "dd/mm/yyyy", "Agreed end of the backtest period."),
    (6, "Field extracted", "PX_LAST", "@", "Closing price (settlement price for futures)."),
]
p["A3"] = "Extraction parameters"
p["A3"].font = f_bold
for r, lab, val, fmt, note in rows:
    p.cell(r, 1, lab).font = f_base
    c = p.cell(r, 2, val)
    c.font, c.fill, c.number_format, c.border = f_input, fill_in, fmt, box
    p.cell(r, 3, note).font = f_note

p["A8"] = "To record at the terminal (mandatory)"
p["A8"].font = f_bold
notes = [
    (9, "Generic roll method", "E.g. 5 days before expiry, or With Active Futures"),
    (10, "Price adjustment type", "E.g. Ratio (expected), Difference, or None"),
    (11, "Extraction date", "E.g. 13/10/2026"),
    (12, "Terminal / login used", "E.g. SKEMA trading room, desk 4"),
]
for r, lab, ex in notes:
    p.cell(r, 1, lab).font = f_base
    c = p.cell(r, 2)
    c.font, c.fill, c.border = f_input, fill_in, box
    p.cell(r, 3, ex).font = f_note
p["B11"].number_format = "dd/mm/yyyy"
p["B10"].comment = Comment(
    "An unadjusted generic contains an artificial jump at every contract roll. "
    "It must be ratio back-adjusted. Setting: GFUT <GO>, Generic Rolls tab; "
    "BDH uses this setting by default.",
    "Claude")

p["A14"] = "How to use"
p["A14"].font = f_bold
steps = [
    "1. Before opening the file: GFUT <GO> → Generic Rolls → Adjust = RATIO and roll BEFORE expiry (Days = 5 or With Active Futures, never at expiry). Fill in B9 and B10. If the setting is locked on a shared login, extract anyway: the Plan B block (2nd generics) lets the returns be rebuilt.",
    "2. Tickers sheet: tickers in column C can be edited; every formula in the Prices sheet follows automatically.",
    "3. Open the file on the Bloomberg terminal, then Bloomberg ribbon > Refresh Workbook. If #NAME? appears: Ctrl+Alt+F9; if it persists, select the formula cell, press F2 then Enter.",
    "4. Wait until every '#N/A Requesting Data...' is gone. The Checks sheet must show 'Yes' in C3 and 'Adjusted' in C6. Read the Status column row by row.",
    "5. Manual check: compare the last value of ES1 Index and CL1 Comdty with the GP screen. A gap means a wrong roll or field setting.",
    "6. Freeze: File > Save As 'Extract_Trend_YYYYMMDD_values.xlsx'. In THAT copy, Tickers then Prices sheets: Ctrl+A, Ctrl+C, Ctrl+Alt+V > Values. Formats are already in place and stay. Leave Checks as formulas: they also work on values.",
    "7. Check the freeze: Ctrl+F, search 'BDH' then 'BDP' in formulas (Options > Look in: Formulas). Zero results expected.",
    "8. Put the '_values' file in data/. Keep the file with formulas as a reusable template.",
]
for i, s in enumerate(steps):
    c = p.cell(15 + i, 1, s)
    c.font, c.alignment = f_base, wrap
    p.merge_cells(start_row=15 + i, start_column=1, end_row=15 + i, end_column=3)
    p.row_dimensions[15 + i].height = 30

p["A24"] = f"Quota: {N} tickers × 1 field in BDH ≈ {N} hits, far below the daily limit (500,000). Every recalculation uses quota again: freeze as soon as the checks are green."
p["A24"].font = f_note
p.merge_cells("A24:C24")
p["A24"].alignment = wrap
p.row_dimensions[24].height = 26

# ---------------------------------------------------------------- Tickers
t = wb.create_sheet("Tickers")
t.sheet_view.showGridLines = False
t.freeze_panes = "A5"
widths = [5, 26, 18, 52, 13, 34, 10, 14, 14]
for i, w in enumerate(widths, 1):
    t.column_dimensions[get_column_letter(i)].width = w
t["A1"] = "Backtest universe"
t["A1"].font = f_title
t["A2"] = "Column C (blue): editable ticker. Columns F to H: terminal check via BDP (name, currency, contract size)."
t["A2"].font = f_note
hdr = ["#", "Block", "Ticker", "Description", "Status", "NAME (BDP)", "CRNCY (BDP)", "FUT_CONT_SIZE (BDP)", "Columns in Prices"]
for j, h in enumerate(hdr, 1):
    c = t.cell(4, j, h)
    c.font, c.fill, c.border = f_h, fill_h, box
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
t.row_dimensions[4].height = 28

for i, (bloc, tick, desc, stat, fut) in enumerate(TICKERS):
    r = 5 + i
    dcol = get_column_letter(1 + 3 * i)
    vcol = get_column_letter(2 + 3 * i)
    vals = [i + 1, bloc, tick, desc, stat,
            f"=BDP(C{r},\"NAME\")",
            f"=BDP(C{r},\"CRNCY\")",
            f"=BDP(C{r},\"FUT_CONT_SIZE\")" if fut else "n/a",
            f"{dcol}:{vcol}"]
    for j, v in enumerate(vals, 1):
        c = t.cell(r, j, v)
        c.font, c.border = f_base, box
    t.cell(r, 3).font = f_input
    if stat in ("To check", "To complete"):
        for j in (3, 5):
            t.cell(r, j).fill = fill_in
t.cell(5 + N + 1, 2, "Ratio adjustment required for every futures row — see Parameters!B10.").font = f_note

# ---------------------------------------------------------------- Prices
x = wb.create_sheet("Prices")
x.freeze_panes = "A4"
date_fmt = "dd/mm/yyyy"
num_fmt = "#,##0.0000"
for i in range(N):
    dcol_i, vcol_i, scol_i = 1 + 3 * i, 2 + 3 * i, 3 + 3 * i
    dcol, vcol = get_column_letter(dcol_i), get_column_letter(vcol_i)
    r_t = 5 + i
    c = x.cell(1, dcol_i, f"=Tickers!$C${r_t}")
    c.font, c.fill = f_link_b, fill_sub
    x.cell(1, vcol_i).fill = fill_sub
    c = x.cell(2, dcol_i, f"=Tickers!$D${r_t}")
    c.font = Font(name=F, size=8, color="008000")
    for col, lab in ((dcol_i, "Date"), (vcol_i, "PX_LAST")):
        c = x.cell(3, col, lab)
        c.font, c.fill, c.border = f_h, fill_h, box
    x.cell(FIRST, dcol_i,
           f"=BDH({dcol}$1,Parameters!$B$6,Parameters!$B$4,Parameters!$B$5,"
           f"\"Dir=V\",\"Dts=S\",\"Sort=A\",\"Days=T\",\"Per=D\")").font = f_base
    for r in range(FIRST, LAST + 1):
        a = x.cell(r, dcol_i)
        a.number_format = date_fmt
        a.font = f_base
        b = x.cell(r, vcol_i)
        b.number_format = num_fmt
        b.font = f_base
    x.column_dimensions[dcol].width = 11
    x.column_dimensions[vcol].width = 13
    x.column_dimensions[get_column_letter(scol_i)].width = 2

# ---------------------------------------------------------------- Checks
k = wb.create_sheet("Checks")
k.sheet_view.showGridLines = False
k.freeze_panes = "A8"
kw = [5, 16, 26, 11, 13, 13, 10, 13, 34]
for i, w in enumerate(kw, 1):
    k.column_dimensions[get_column_letter(i)].width = w
k["A1"] = "Checks before freezing"
k["A1"].font = f_title
k["A3"] = "Ready to freeze?"
k["A3"].font = f_bold
k.merge_cells("A3:B3")
first_r, last_r = 8, 8 + N - 1
k["C3"] = f'=IF(AND(SUM(G{first_r}:G{last_r})=0,LEFT(C6,3)<>"NOT"),"Yes","No — see Status and adjustment test")'
k["C3"].font = f_bold
k["A4"] = "Suspicious daily jump threshold"
k["A4"].font = f_base
k.merge_cells("A4:B4")
k["C4"] = 0.25
k["C4"].font, k["C4"].fill, k["C4"].number_format = f_input, fill_in, "0%"
k["D4"] = ("Assumption: a daily move > 25% on an adjusted future probably signals an unadjusted roll "
           "(known exceptions: CL1 in April 2020, negative price; silver on 30/01/2026 and copper on 31/07/2025 are real moves). Not applied to short rates.")
k["D4"].font = f_note
k["A5"] = "Minimum number of points expected"
k["A5"].font = f_base
k.merge_cells("A5:B5")
k["C5"] = 5000
k["C5"].font, k["C5"].fill, k["C5"].number_format = f_input, fill_in, "#,##0"
k["D5"] = "Assumption: ≈ 250 sessions × 20 years. Below that, history is short (information, not necessarily an error)."
k["D5"].font = f_note
k["A6"] = "Futures adjustment test"
k["A6"].font = f_bold
k.merge_cells("A6:B6")
k["C6"] = '=IF(Prices!A4<>DATE(2000,1,3),"Different start date — test not applicable",IF(ABS(Prices!B4/1466.75-1)<0.005,"NOT ADJUSTED — redo with GFUT set to Ratio","Adjusted"))'
k["C6"].font = f_bold
k["D6"] = ("Raw ES1 was 1,466.75 on 03/01/2000 (source: unadjusted extraction of 13/10/2026). "
           "A ratio back-adjusted series shows a different level.")
k["D6"].font = f_note
k.conditional_formatting.add("C6", FormulaRule(formula=['$C$6="Adjusted"'], fill=PatternFill("solid", fgColor="C6EFCE")))
k.conditional_formatting.add("C6", FormulaRule(formula=['LEFT($C$6,3)="NOT"'], fill=PatternFill("solid", fgColor="F8CBAD")))

khdr = ["#", "Ticker", "Block", "Points", "First date", "Last date", "Errors", "Largest daily move", "Status"]
for j, h in enumerate(khdr, 1):
    c = k.cell(7, j, h)
    c.font, c.fill, c.border = f_h, fill_h, box
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
k.row_dimensions[7].height = 28

for i in range(N):
    r = first_r + i
    dcol = get_column_letter(1 + 3 * i)
    vcol = get_column_letter(2 + 3 * i)
    rng_v = f"Prices!{vcol}{FIRST}:{vcol}{LAST}"
    rng_v_prev = f"Prices!{vcol}{FIRST - 1}:{vcol}{LAST - 1}"
    rng_d = f"Prices!{dcol}{FIRST}:{dcol}{LAST}"
    rng_all = f"Prices!{dcol}{FIRST}:{vcol}{LAST}"
    cells = [
        i + 1,
        f"=Tickers!C{5 + i}",
        f"=Tickers!B{5 + i}",
        f"=COUNT({rng_v})",
        f"=IF(D{r}=0,\"\",MIN({rng_d}))",
        f"=IF(D{r}=0,\"\",MAX({rng_d}))",
        f"=SUMPRODUCT(ISERROR({rng_all})+ISTEXT({rng_all}))",
        f"=IF(D{r}<2,\"\",_xlfn.AGGREGATE(14,6,ABS({rng_v}/{rng_v_prev}-1)/(ISNUMBER({rng_v})*ISNUMBER({rng_v_prev})),1))",
        (f"=IF(B{r}=\"\",\"Missing ticker\",IF(G{r}>0,\"Errors or data pending\","
         f"IF(D{r}=0,\"No data\","
         f"IF(F{r}<Parameters!$B$5-7,\"Series stops before the end date\","
         f"IF(AND(C{r}<>\"Short rates\",LEFT(C{r},6)<>\"Plan B\",H{r}>$C$4),\"Suspicious jump: unadjusted roll?\","
         f"IF(D{r}<$C$5,\"OK — short history\",\"OK\"))))))"),
    ]
    fmts = ["0", "@", "@", "#,##0", "dd/mm/yyyy", "dd/mm/yyyy", "0", "0.0%", "@"]
    for j, (v, fm) in enumerate(zip(cells, fmts), 1):
        c = k.cell(r, j, v)
        c.font, c.border, c.number_format = f_base, box, fm
    k.cell(r, 2).font = f_link
    k.cell(r, 3).font = f_link

red = PatternFill("solid", fgColor="F8CBAD")
green = PatternFill("solid", fgColor="C6EFCE")
k.conditional_formatting.add(f"I{first_r}:I{last_r}",
                             FormulaRule(formula=[f'LEFT($I{first_r},2)="OK"'], fill=green))
k.conditional_formatting.add(f"I{first_r}:I{last_r}",
                             FormulaRule(formula=[f'LEFT($I{first_r},2)<>"OK"'], fill=red))
k.conditional_formatting.add("C3", FormulaRule(formula=['$C$3="Yes"'], fill=green))
k.conditional_formatting.add("C3", FormulaRule(formula=['$C$3<>"Yes"'], fill=red))
k.cell(last_r + 2, 1, "Errors also counts text such as '#N/A Requesting Data...' or '#N/A N/A' returned by Bloomberg.").font = f_note
k.cell(last_r + 3, 1, "Largest daily move = largest absolute change from one session to the next in the series.").font = f_note

wb.active = 0
wb.save(OUT)
print("saved", OUT)
