"""Load a Bloomberg extract that has been frozen as values (see templates/bloomberg).

Expected layout of the price sheet ("Prices" in the English template, "Prix" in the French one):
one block of 3 columns per ticker (date, PX_LAST, empty spacer). Row 1 = ticker, rows 2-3 = labels,
data from row 4 onwards.
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

PRICE_SHEETS = ("Prices", "Prix")


def load_bloomberg_extract(path: str | Path) -> dict[str, pd.Series]:
    """Return {ticker: price series indexed by date}, de-duplicated and sorted."""
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet = next((s for s in PRICE_SHEETS if s in wb.sheetnames), None)
    if sheet is None:
        raise ValueError(f"No price sheet found (expected one of {PRICE_SHEETS}) in {path}")
    rows = list(wb[sheet].iter_rows(values_only=True))
    header = rows[0]
    out = {}
    for j in range(0, len(header), 3):
        ticker = header[j]
        if not ticker:
            continue
        pairs = [(r[j], r[j + 1]) for r in rows[3:]
                 if r[j] is not None and isinstance(r[j + 1], (int, float))]
        if not pairs:
            continue
        s = pd.Series([v for _, v in pairs],
                      index=pd.to_datetime([d for d, _ in pairs]), name=ticker)
        out[ticker] = s[~s.index.duplicated()].sort_index()
    return out


def excess_returns(prices: pd.Series, rf_annual_pct: pd.Series, days_per_year: int = 261) -> pd.Series:
    """Daily return minus that day's risk-free rate (annual rate in %, e.g. 3-month Euribor)."""
    r = prices.pct_change().dropna()
    rf = (rf_annual_pct / 100 / days_per_year).reindex(r.index).ffill().fillna(0)
    return r - rf
